#!/usr/bin/env python3
"""Probes for ingest_uscis_rfe.py: names, discovery, the three formats, every
reconciliation refusal, and the whole run against real in-memory SQLite.

    python3 scripts/test_uscis_rfe.py

NO NETWORK and no production database. The workbooks are built here with
openpyxl in the two shapes USCIS has published (the single-year edition with
a Pending column, and the multi-year edition without one, whose name starts
a year before its first row), every total computed by the builder, so a
refusal probe breaks a parsed value and never the file. The PDF editions are
probed through their extracted text, which is what the parser reads.
"""
from __future__ import annotations

import io
import os
import pathlib
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ingest_uscis_rfe as rfe  # noqa: E402
from lib_db import cell  # noqa: E402
from lib_sqlite_shim import SqliteTurso  # noqa: E402

FAILURES: list[str] = []
FY_MONTHS = ["October", "November", "December", "January", "February", "March",
             "April", "May", "June", "July", "August", "September"]


def check(name: str, cond: bool, detail: str = "") -> None:
    print(f"  {'PASS' if cond else 'FAIL'}  {name}{(' - ' + detail) if detail and not cond else ''}")
    if not cond:
        FAILURES.append(name)


def refuses(name: str, fn, needle: str) -> None:
    try:
        fn()
        check(name, False, "it was accepted")
    except rfe.Refusal as exc:
        check(name, needle in str(exc), str(exc)[:160])


# ---------------------------------------------------------------------------
# Builders
# ---------------------------------------------------------------------------

def month_values(seed: int, i: int) -> dict:
    """Counts that satisfy every identity the file is checked against."""
    approved = 100 + seed * 7 + i * 3
    denied = 5 + (seed + i) % 4
    rfe_c = 10 + (seed * 3 + i) % 9
    return dict(received=120 + seed + i, approved=approved, denied=denied,
                completions=approved + denied, rfe=rfe_c, rfe_ok=rfe_c - 2, pending=900 - i * 10)


def sheet_rows(title: str, span: str, fys: list[int], months_last: int, seed: int,
               pending: bool, year_totals: bool) -> list[list]:
    head = ["Fiscal Year", "Month", "Petitions Received 1", "Initially Approved 2", "Initially Denied3",
            "Total Completions4"]
    if pending:
        head.append("Pending5")
    head += ["Approved (%)6", "Completions with RFE7", "Completions with RFE (%)8",
             "Approved with RFE9", "Approved with RFE (%)10"]

    def line(first, month, v):
        row = [first, month, v["received"], v["approved"], v["denied"], v["completions"]]
        if pending:
            row.append(v["pending"])
        row += [round(v["approved"] / v["completions"], 3), v["rfe"], round(v["rfe"] / v["completions"], 3),
                v["rfe_ok"], round(v["rfe_ok"] / v["completions"], 3)]
        return row

    body, every = [], []
    for fy in fys:
        n = months_last if fy == fys[-1] else 12
        vals = [month_values(seed + fy, i) for i in range(n)]
        every += vals
        for i, v in enumerate(vals):
            body.append(line(str(fy) if i == 0 else None, FY_MONTHS[i], v))
        if year_totals:
            t = {k: sum(v[k] for v in vals) for k in vals[0]}
            t["pending"] = vals[-1]["pending"]
            body.append(line(f"{fy} Total", None, t))
    total = {k: sum(v[k] for v in every) for k in every[0]}
    total["pending"] = every[-1]["pending"]
    return [[title], ["By Fiscal Year, Month, and Case Status"], [span], ["Period", None, "Petitions by Case Status"],
            head, line("TOTAL", None, total), *body,
            ["1) The number of new petitions received."], ["Sources:"]]


def build_xlsx(path: pathlib.Path, tables: list[tuple[str, str]], span: str, fys: list[int],
               months_last: int, pending: bool, year_totals: bool, notes: bool = True) -> None:
    import openpyxl  # only the builder needs it

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for i, (sheet, title) in enumerate(tables):
        ws = wb.create_sheet(sheet)
        for row in sheet_rows(title, span, fys, months_last, seed=i * 11, pending=pending,
                              year_totals=year_totals):
            ws.append(row)
    if notes:
        ws = wb.create_sheet("i129_notes")
        ws.append(["References:"])
        ws.append(["1) The number of new petitions received."])
    wb.save(path)


H1B = "Form I-129 Petition for a Nonimmigrant Worker Specialty Occupations (H-1B)"
L1A = "Form I-129 Petition for a Nonimmigrant Worker, Intracompany Transferee Executive or Manager (L-1A) "
L1 = ("Form I-129 Petition for a Nonimmigrant Worker, Intracompany Transferee  (L-1), "
      "Not Assigned L-1A or L-1B")
BLANKET = "Form I-129 Petition for a Nonimmigrant Worker, Blanket L Petition"


def parse_file(path: pathlib.Path) -> rfe.Parsed:
    item = rfe.identify(path.name, str(path))
    assert item, path.name
    return rfe.parse_blob(item, path.read_bytes())


def copy_of(p: rfe.Parsed) -> rfe.Parsed:
    tables = [rfe.Table(t.visa, t.start, t.end, [dict(r) for r in t.months],
                        {k: dict(v) for k, v in t.year_totals.items()},
                        dict(t.grand) if t.grand else None, t.has_pending) for t in p.tables]
    return rfe.Parsed(p.listed, tables)


# ---------------------------------------------------------------------------
# PDF text, as pdfplumber extracts it
# ---------------------------------------------------------------------------

PDF_HEAD = ["Period Petitions by Case Status",
            "Petitions Initially Initially Total Approved Completions Completions with Approved Approved with RFE",
            "Fiscal Year Month Received 1 Approved 2 Denied3 Completions 4 (%)5 with RFE6 RFE (%)7 with RFE8 (%)9"]


def pdf_line(first: str, v: dict) -> str:
    c = v["completions"]
    return (f"{first} {v['received']:,} {v['approved']:,} {v['denied']:,} {c:,} {v['approved'] / c:.1%} "
            f"{v['rfe']:,} {v['rfe'] / c:.1%} {v['rfe_ok']:,} {v['rfe_ok'] / v['rfe']:.1%}")


def pdf_pages_for(fys: list[int]) -> list[str]:
    title = ["I-129 - Petition for a Nonimmigrant Worker", "Specialty Occupations (H-1B)",
             "by Fiscal Year, Month, and Case Status:", f"October 1, {fys[0] - 1} - September 30, {fys[-1]}"]
    lines, every = [*title, *PDF_HEAD], []
    for fy in fys:
        vals = [month_values(fy, i) for i in range(12)]
        every += vals
        for i, v in enumerate(vals):
            lines.append(pdf_line(f"{fy} {FY_MONTHS[i]}" if i == 0 else FY_MONTHS[i], v))
        lines.append(pdf_line(f"{fy} Total", {k: sum(x[k] for x in vals) for k in vals[0]}))
    lines.append(pdf_line("Grand Total", {k: sum(x[k] for x in every) for k in every[0]}))
    quarter = ["I-129 - Petition for a Nonimmigrant Worker", "Specialty Occupations (H-1B)",
               "by Fiscal Year, Quarter, and Case Status:", title[3],
               "Fiscal Year Quarter Received 1 Approved 2 Denied3 Completions 4 (%)5 with RFE6 RFE (%)7 with RFE8 (%)9",
               "2015 Q1. Oct. - Dec. 1,375 1,147 9 1,156 99.2% 136 11.8% 127 93.4%"]
    # The month table runs onto a second page with its headings repeated.
    half = 4 + len(PDF_HEAD) + 13
    return ["\n".join(lines[:half]), "\n".join(PDF_HEAD + lines[half:]), "\n".join(quarter)]


# ---------------------------------------------------------------------------
# Probes
# ---------------------------------------------------------------------------

def main() -> int:
    print("ingest_uscis_rfe probes\n")
    tmp = pathlib.Path(tempfile.mkdtemp(prefix="rfe-"))

    # -- names and discovery -------------------------------------------------
    a = rfe.identify("i129_quarterly_request_for_evidence_fy2026_q3_v1.xlsx")
    b = rfe.identify("i129_quarterly_request_for_evidence_fy2017_fy2024_q4.xlsx")
    c = rfe.identify("I129_Quarterly_Request_for_Evidence_FY2015_FY2021_Q4.pdf")
    check("names: single-year, multi-year and the old capitalised PDF all read",
          (a.last_fy, a.quarter, a.version, a.fmt) == (2026, 3, 1, "xlsx")
          and (b.named_first, b.last_fy, b.quarter) == (2017, 2024, 4)
          and (c.named_first, c.last_fy, c.fmt) == (2015, 2021, "pdf"))
    check("names: other files on the page are not ours",
          rfe.identify("quarterly_all_forms_fy2026_q3.xlsx") is None
          and rfe.identify("i129_quarterly_request_for_evidence_fy2026_q3.docx") is None)
    v2 = rfe.identify("i129_quarterly_request_for_evidence_fy2026_q3_v2.xlsx")
    csv_q = rfe.identify("I129_Quarterly_Request_for_Evidence_FY2016_FY2022_Q4.csv")
    pdf_q = rfe.identify("I129_Quarterly_Request_for_Evidence_FY2016_FY2022_Q4.pdf")
    check("rank: a re-issued _v2 beats _v1, a later quarter beats both, a spreadsheet beats a CSV beats a PDF",
          v2.rank > a.rank > b.rank and csv_q.rank > pdf_q.rank)
    check("quarter end: Q3 ends June 30, Q1 the December before",
          (a.end, rfe.identify("i129_quarterly_request_for_evidence_fy2026_q1_v1.xlsx").end)
          == ("2026-06-30", "2025-12-31"))
    html = ('<a href="/sites/default/files/document/data/i129_quarterly_request_for_evidence_fy2026_q3_v1.xlsx">x</a>'
            '<a href="/sites/default/files/document/reports/I129_Quarterly_Request_for_Evidence_FY2015_FY2021_Q4.pdf">y</a>'
            '<a href="/sites/default/files/document/data/quarterly_all_forms_fy2026_q3.xlsx">z</a>'
            '<a href="https://www.uscis.gov/x/i129_quarterly_request_for_evidence_fy2026_q3_v1.xlsx">dup</a>')
    found = {x.name: x.url for x in rfe.discover(html)}
    check("discover: both RFE files found, the all-forms file left alone, a repeat counted once",
          len(found) == 2 and found["i129_quarterly_request_for_evidence_fy2026_q3_v1.xlsx"].startswith(
              "https://www.uscis.gov/sites/default/files/document/data/"))
    check("discover: the listing's later pages are the page's own control",
          rfe.page_url(0) == rfe.DATA_PAGE and rfe.page_url(3) == rfe.DATA_PAGE + "&page=3")
    check("plan: without --history only the newest file is read",
          [x.name for x in rfe.plan([b, a, c], history=False)] == [a.name]
          and [x.name for x in rfe.plan([c, a, b], history=True)] == [a.name, b.name, c.name])

    # -- cells -------------------------------------------------------------------
    check("cells: thousands, a dash, a percent string and a stored fraction",
          rfe.count("2,542,134", "") == 2542134 and rfe.count("-", "") == 0 and rfe.count("", "") is None
          and abs(rfe.share("8.5%", "") - 0.085) < 1e-12 and abs(rfe.share("8.5000000000000006E-2", "") - 0.085) < 1e-12)
    refuses("cells: a fraction where a count belongs is refused", lambda: rfe.count("12.5", "x"), "whole count")
    check("headings: footnote digits come off", rfe.strip_note("Initially Denied3") == "Initially Denied"
          and rfe.strip_note("Completions with RFE (%)8") == "Completions with RFE (%)"
          and rfe.strip_note("Total Completions 4") == "Total Completions")
    check("visas: every title maps, and the unassigned L-1 is not the L-1A",
          [rfe.visa_of(t) for t in (H1B, L1A, L1, BLANKET,
                                    "Intracompany Transferee Executive or Manager (L-1A) Intercompany Transferee "
                                    "Specialized Knowledge (L-1B)")]
          == ["h1b", "l1a", "l1_unassigned", "l1_blanket", "l1"])

    # -- the single-year edition, with Pending ---------------------------------
    one = tmp / "i129_quarterly_request_for_evidence_fy2026_q2_v1.xlsx"
    build_xlsx(one, [("i129_FY26q2_h1b", H1B), ("i129_FY26q2_l1a", L1A), ("i129_FY26q2_Blanket_L", BLANKET)],
               "October 1, 2025 - March 31, 2026", [2026], 6, pending=True, year_totals=False)
    p1 = parse_file(one)
    h = next(t for t in p1.tables if t.visa == "h1b")
    check("single-year: three tables, the notes sheet skipped", [t.visa for t in p1.tables] == ["h1b", "l1a", "l1_blanket"])
    check("single-year: October is calendar month 10 of FY2026, March month 3",
          (h.months[0]["fy"], h.months[0]["month"], h.months[-1]["month"]) == (2026, 10, 3))
    check("single-year: pending is read", h.has_pending and h.months[0]["pending"] == 900)
    check("single-year: six months per visa", p1.fys() == [2026] and len(p1.rows_for(2026)) == 18)

    # -- the multi-year edition, no Pending, named a year early ----------------
    many = tmp / "i129_quarterly_request_for_evidence_fy2016_fy2018_q4.xlsx"
    build_xlsx(many, [("i129_FY18q4_h1b", H1B), ("i129_FY18q4_L1A", L1A), ("i129_FY18q4_L1", L1)],
               "October 1, 2016 - September 30, 2018", [2017, 2018], 12, pending=False, year_totals=True)
    pm = parse_file(many)
    check("multi-year: the years come from the rows, not the name", pm.fys() == [2017, 2018])
    check("multi-year: no Pending column reads as None", all(r["pending"] is None for r in pm.rows_for(2017)))
    check("multi-year: the unassigned L-1 table is left out", [t.visa for t in pm.tables] == ["h1b", "l1a"])
    check("multi-year: year totals are read", sorted(next(t for t in pm.tables if t.visa == "h1b").year_totals) == [2017, 2018])

    # -- refusals ----------------------------------------------------------------
    def mutated(fn) -> None:
        bad = copy_of(pm)
        fn(bad)
        rfe.reconcile(bad.listed, bad.tables)

    def bump_month(bad, field):
        t = bad.tables[0]
        t.months[4][field] += 1
        if field in ("approved", "denied"):
            t.months[4]["completions"] += 1

    def shift_between_years(bad):
        # One RFE moved from an FY2018 month to an FY2017 month: the TOTAL row
        # still adds up, so only the per-year totals can catch it.
        t = bad.tables[0]
        t.months[4]["rfe_completions"] += 1
        t.months[16]["rfe_completions"] -= 1

    refuses("refuse: a month that no longer sums to its fiscal year's total",
            lambda: mutated(shift_between_years), "FY2017: rfe_completions months sum to")
    refuses("refuse: a month that no longer sums to the TOTAL row",
            lambda: mutated(lambda bad: bump_month(bad, "rfe_completions")), "the TOTAL row says")
    refuses("refuse: approved plus denied that isn't completions",
            lambda: mutated(lambda bad: bad.tables[0].months[2].__setitem__("denied", bad.tables[0].months[2]["denied"] + 1)),
            "is not completions")
    refuses("refuse: a printed RFE share 0.002 off the counts",
            lambda: mutated(lambda bad: bad.tables[0].months[1].__setitem__("rfe_share", bad.tables[0].months[1]["rfe_share"] + 0.002)),
            "printed RFE share")
    within = copy_of(pm)
    within.tables[0].months[1]["rfe_share"] += 0.001
    check("accept: a share within the 0.0015 rounding allowance", not rfe.problems_in(within.listed, within.tables))
    refuses("refuse: a month missing from the run",
            lambda: mutated(lambda bad: bad.tables[1].months.pop(7)), "without a gap")
    refuses("refuse: a table that ends on another quarter than the name's",
            lambda: mutated(lambda bad: setattr(bad.tables[0], "end", "2018-06-30")), "named for")
    refuses("refuse: a TOTAL row the months don't reach",
            lambda: mutated(lambda bad: bad.tables[1].grand.__setitem__("received", bad.tables[1].grand["received"] + 5)),
            "TOTAL row says")
    refuses("refuse: a file with no H-1B table",
            lambda: mutated(lambda bad: bad.tables.pop(0)), "no H-1B table")
    single = copy_of(p1)
    single.tables[0].grand["pending"] += 1
    refuses("refuse: a TOTAL pending that isn't the newest month's",
            lambda: rfe.reconcile(single.listed, single.tables), "pending")
    every = copy_of(pm)
    every.tables[0].months[0]["denied"] += 1
    every.tables[1].months[0]["rfe_share"] = 0.9
    check("one refusal names every problem, not only the first",
          len(rfe.problems_in(every.listed, every.tables)) >= 3)
    stranger = tmp / "i129_quarterly_request_for_evidence_fy2026_q1_v1.xlsx"
    build_xlsx(stranger, [("h1b", H1B), ("x", "Form I-129 Petition for a Nonimmigrant Worker, Someone New (Q-9)")],
               "October 1, 2025 - December 31, 2025", [2026], 3, pending=True, year_totals=False)
    refuses("refuse: a table under a title no visa key matches", lambda: parse_file(stranger), "no visa key")

    # -- CSV -----------------------------------------------------------------------
    rows = sheet_rows(H1B, '"October 1, 2016 - September 30, 2017"', [2017], 12, 3, pending=False, year_totals=True)
    out = io.StringIO()
    for r in rows:
        cells = []
        for x in r:
            if x is None:
                cells.append("")
            elif isinstance(x, float):
                cells.append(f"{x:.1%}")
            elif isinstance(x, int):
                cells.append(f'"{x:,}"')
            else:
                cells.append(f'"{x}"' if "," in x and not x.startswith('"') else x)
        out.write(",".join(cells) + "\n")
        out.write("This cell left intentionally blank,This cell left intentionally blank\n")
    blocks = rfe.csv_blocks(("﻿,,,\n" + out.getvalue()).encode())
    t = rfe.parse_block(*blocks[0], where="csv")
    item = rfe.identify("I129_Quarterly_Request_for_Evidence_FY2016_FY2017_Q4.csv")
    check("csv: one table, twelve months, the blank-cell rows ignored, and it reconciles",
          len(blocks) == 1 and t.visa == "h1b" and len(t.months) == 12 and not rfe.problems_in(item, [t]),
          str(rfe.problems_in(item, [t]))[:200] if t else "")

    # -- PDF text --------------------------------------------------------------------
    pages = pdf_pages_for([2015, 2016])
    tables = rfe.pdf_tables(pages, "pdf")
    pdf_item = rfe.identify("I129_Quarterly_Request_for_Evidence_FY2015_FY2016_Q4.pdf")
    check("pdf: the date range under the title isn't read as an October row",
          len(tables) == 1 and len(tables[0].months) == 24 and tables[0].months[0]["fy"] == 2015)
    check("pdf: the quarter table is skipped and the month table reconciles across its page break",
          not rfe.problems_in(pdf_item, tables), str(rfe.problems_in(pdf_item, tables))[:200])
    check("pdf: a combined L-1A and L-1B table is its own key",
          rfe.visa_of(" ".join(["I-129 - Petition for a Nonimmigrant Worker",
                                "Intracompany Transferee Executive or Manager (L-1A)",
                                "Intercompany Transferee Specialized Knowledge (L-1B)"])) == "l1")
    changed = [pages[0].replace("Completions 4 (%)5", "Completions 4 Pending (%)5"), *pages[1:]]
    refuses("pdf: changed column headings are refused, not guessed at", lambda: rfe.pdf_tables(changed, "pdf"),
            "headings changed")
    short = [pages[0].replace(pdf_line("November", month_values(2015, 1)),
                              pdf_line("November", month_values(2015, 1)).rsplit(" ", 1)[0]), *pages[1:]]
    refuses("pdf: a row with a figure missing is refused", lambda: rfe.pdf_tables(short, "pdf"), "expected 9 figures")

    check("pdf: the later layout's garbled codes are rebuilt before matching",
          rfe.visa_of(rfe.clean_title("Form I--129 Petition for Nonimmigrant Worker Specialty Occupations (H -1B)")) == "h1b"
          and rfe.visa_of(rfe.clean_title("Form I--129 Petition, Temporary Agricultural Worker (-H 2A)")) == "h2a")
    later = ["\n".join([
        "Form I--129 Petition for Nonimmigrant Worker Specialty Occupations (H -1B)",
        "By Fiscal Year, Month, and Case Status", "October 1, 2014 - September 30, 2015", "Period Petitions by Case Status",
        "Fiscal Year Month Petitions Received1 Initially Approved2 Initially Denied3 Total Completions4 Approved (%)5 "
        "Completions with RFE6 Completions with RFE (%)7 Approved with RFE8 Approved with RFE (%)9",
        pdf_line("TOTAL", {k: sum(month_values(2015, i)[k] for i in range(12)) for k in month_values(2015, 0)}),
        *[pdf_line(f"2015 {FY_MONTHS[i]}" if i == 0 else FY_MONTHS[i], month_values(2015, i)) for i in range(12)],
    ])]
    later_tables = rfe.pdf_tables(later, "pdf")
    check("pdf: the later one-line headings are read, the TOTAL row at the top",
          len(later_tables) == 1 and len(later_tables[0].months) == 12 and later_tables[0].grand is not None
          and not rfe.problems_in(rfe.identify("I129_Quarterly_Request_for_Evidence_FY2015_FY2015_Q4.pdf"), later_tables))
    check("span: the fiscal years a date range covers",
          (rfe.fiscal_year_of("2016-10-01"), rfe.fiscal_year_of("2023-06-30")) == (2017, 2023))

    # -- the whole run, against real SQLite ---------------------------------------
    db = SqliteTurso()
    rfe.Database = lambda: db  # type: ignore[assignment]
    q1 = tmp / "q1" / "i129_quarterly_request_for_evidence_fy2026_q1_v1.xlsx"
    q1.parent.mkdir()
    build_xlsx(q1, [("h1b", H1B), ("l1a", L1A)], "October 1, 2025 - December 31, 2025", [2026], 3,
               pending=True, year_totals=False)
    rc = rfe.main(["--file", str(q1)])
    n = db.scalar("SELECT count(*) FROM uscis_i129_rfe")
    check("run: exits 0 and stores every parsed row", rc == 0 and int(n) == 6, f"rc {rc}, rows {n}")
    rc = rfe.main(["--file", str(one), "--file", str(many)])
    by_fy = {int(cell(r[0])): (int(cell(r[1])), cell(r[2])) for r in db.execute(
        "SELECT fy, count(*), max(source_file) FROM uscis_i129_rfe GROUP BY fy")["response"]["result"]["rows"]}
    check("run: the newer quarter replaces FY2026, and the multi-year file adds FY2017 and FY2018",
          rc == 0 and by_fy == {2017: (24, many.name), 2018: (24, many.name), 2026: (18, one.name)}, str(by_fy))
    rc = rfe.main(["--file", str(q1)])
    still = db.scalar("SELECT max(source_file) FROM uscis_i129_rfe WHERE fy = 2026")
    check("run: an older file never overwrites a year held from a newer one",
          rc == 0 and still == one.name and int(db.scalar("SELECT count(*) FROM uscis_i129_rfe WHERE fy = 2026")) == 18)
    # Older editions of years already read: one that doesn't add up, and one
    # that can't even be parsed (as the later PDFs' split figures can't). Each
    # is set aside on its date range, so neither can refuse the run.
    stale = tmp / "stale" / "i129_quarterly_request_for_evidence_fy2016_fy2018_q3.xlsx"
    stale.parent.mkdir()
    build_xlsx(stale, [("h1b", H1B)], "October 1, 2016 - June 30, 2018", [2017, 2018], 9,
               pending=False, year_totals=True)
    import openpyxl
    wb = openpyxl.load_workbook(stale)
    wb["h1b"]["C6"] = 1  # the TOTAL row's received, now wrong
    wb.save(stale)
    unreadable = tmp / "stale" / "i129_quarterly_request_for_evidence_fy2016_fy2018_q2.xlsx"
    build_xlsx(unreadable, [("h1b", H1B), ("x", "Form I-129 Petition for a Nonimmigrant Worker, Someone New (Q-9)")],
               "October 1, 2016 - March 31, 2018", [2017, 2018], 6, pending=False, year_totals=True)
    refuses("an edition that doesn't add up is refused when it's read alone", lambda: parse_file(stale), "TOTAL row says")
    refuses("an edition that can't be parsed is refused when it's read alone", lambda: parse_file(unreadable), "no visa key")
    rc = rfe.main(["--file", str(many), "--file", str(stale), "--file", str(unreadable), "--dry-run"])
    check("history: older editions whose years a newer file gave are set aside on their date range", rc == 0)
    rows_before = int(db.scalar("SELECT count(*) FROM uscis_i129_rfe"))
    rc = rfe.main(["--file", str(one)])
    note = db.scalar("SELECT note FROM ingest_runs ORDER BY rowid DESC LIMIT 1")
    check("run: a file already held is a quiet run that says so",
          rc == 0 and "nothing new" in str(note) and int(db.scalar("SELECT count(*) FROM uscis_i129_rfe")) == rows_before,
          str(note))
    fresh = db.execute("SELECT as_of, max_age_days FROM data_freshness WHERE dataset = 'uscis-rfe'")["response"]["result"]["rows"]
    check("run: freshness is the newest quarter's last day, with the 200-day budget",
          fresh and str(cell(fresh[0][0])) == "2026-03-31" and int(cell(fresh[0][1])) == 200, str(fresh))
    import json
    doc = json.loads(db.scalar("SELECT json FROM perm_docs WHERE key = 'uscis_rfe_summary'"))
    h26 = next(y for y in doc["visas"]["h1b"]["years"] if y["fy"] == 2026)
    want = sum(r["rfe_completions"] for r in p1.rows_for(2026) if r["visa"] == "h1b")
    check("doc: newest quarter, year range, and the H-1B year's RFE count from the stored rows",
          doc["quarterLabel"] == "FY2026 Q2" and doc["asOf"] == "2026-03-31" and (doc["firstFy"], doc["lastFy"]) == (2017, 2026)
          and h26["months"] == 6 and h26["rfeCompletions"] == want and doc["currentFy"] == {"fy": 2026, "months": 6, "complete": False},
          str(doc.get("currentFy")))
    check("doc: every visa carries its label", doc["visas"]["l1a"]["label"] == "L-1A")

    lossy = SqliteTurso()
    real_scalar = lossy.scalar
    lossy.scalar = lambda sql, args=None: int(real_scalar(sql, args) or 0) - 1 if "count(*)" in sql else real_scalar(sql, args)  # type: ignore[assignment]
    rfe.Database = lambda: lossy  # type: ignore[assignment]
    refuses("run: a year whose stored count comes back short is refused", lambda: rfe.main(["--file", str(one)]),
            "stored")

    print(f"\n{len(FAILURES)} failure(s)" if FAILURES else "\nall passed")
    for f in FAILURES:
        print(f"  - {f}")
    return 1 if FAILURES else 0


if __name__ == "__main__":
    sys.exit(main())
