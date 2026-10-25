"""What every state's WARN reader shares: reading values, building a row,
giving it an id, and fetching politely.

ingest_warn.py holds the four states read since September 2026 and the write
path; warn_states.py holds every other state. Both build rows through
`notice()` and `assign_ids()`, so a row from any state has the same shape and
the same id rule.

Every request goes through `get()` or `post()`, which ask the host's
robots.txt first (lib_gov_data.robots_allowed, RFC 9309 longest match) and
refuse a path the host asks automated clients not to read. Nothing here gets
past a bot wall or a challenge page: a state that answers scripts with one is
left out of the reader list, not worked around.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import html
import io
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from html.parser import HTMLParser
from typing import NamedTuple

from lib_gov_data import BROWSER_HEADERS, RobotsDisallowed, fetch as gov_fetch, robots_allowed

# Between two requests to the same state's host in one run.
PAUSE_S = 1.0


def pause(seconds: float = PAUSE_S) -> None:
    time.sleep(seconds)


# --- values ---------------------------------------------------------------

MONTHS = {m: i for i, m in enumerate(
    ["january", "february", "march", "april", "may", "june", "july", "august",
     "september", "october", "november", "december"], start=1)}
MONTHS.update({m[:3]: i for m, i in list(MONTHS.items())})
MONTHS["sept"] = 9

_ISO = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[T ].*)?$")
_MDY = re.compile(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4}|\d{2})$")
_NAMED = re.compile(r"^([A-Za-z]{3,9})\.?\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})$")
_DAY_NAMED = re.compile(r"^(\d{1,2})[\s-]([A-Za-z]{3,9})[\s-](\d{4}|\d{2})$")
_COMPACT = re.compile(r"^(\d{4})(\d{2})(\d{2})$")
# A date somewhere inside a longer string ("4/24/25, 7/3/25", "12/31/2026-1/31/2027").
DATE_TOKEN = re.compile(
    r"\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{2,4}|\d{1,2}-\d{1,2}-\d{2,4}"
    r"|[A-Za-z]{3,9}\.?\s+\d{1,2}(?:st|nd|rd|th)?,?\s+\d{4}")


def _year(y: str) -> int:
    n = int(y)
    if len(y) == 4:
        return n
    # A two-digit year: the century that puts it no more than five years ahead.
    pivot = (dt.date.today().year % 100) + 5
    return 2000 + n if n <= pivot else 1900 + n


def _make(y: int, m: int, d: int) -> str | None:
    try:
        return dt.date(y, m, d).isoformat()
    except ValueError:
        return None


def parse_date(v) -> str | None:
    """One date as ISO, or None. Never guesses past what the cell says.

    Reads a spreadsheet's own date or datetime, ISO (with or without a time),
    month/day/year with a two- or four-digit year and either separator,
    "October 21, 2026", "Oct 21 2026", "21-Oct-2026" and "20261021". Anything
    else, a bare month included, is None: the caller decides what a cell
    without a day means.
    """
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        # A spreadsheet date stored as its serial number (Kentucky's 2019 sheet
        # has 43490 among real dates). Only a plausible range: 1954 to 2119.
        if 20000 <= v <= 80000:
            return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
        return None
    s = re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip().rstrip(".")
    # A time after the date ("10/6/2026 12:00:00 AM") says nothing more.
    s = re.sub(r"\s+\d{1,2}:\d{2}(:\d{2})?\s*([AaPp][Mm])?$", "", s)
    if not s:
        return None
    m = _ISO.match(s)
    if m:
        return _make(int(m[1]), int(m[2]), int(m[3]))
    m = _MDY.match(s)
    if m:
        return _make(_year(m[3]), int(m[1]), int(m[2]))
    m = _NAMED.match(s)
    if m and m[1].lower().rstrip(".") in MONTHS:
        return _make(int(m[3]), MONTHS[m[1].lower()], int(m[2]))
    m = _DAY_NAMED.match(s)
    if m and m[2].lower() in MONTHS:
        return _make(_year(m[3]), MONTHS[m[2].lower()], int(m[1]))
    m = _COMPACT.match(s)
    if m:
        return _make(int(m[1]), int(m[2]), int(m[3]))
    return None


def first_date(v) -> str | None:
    """The first date in a cell that may hold several or a range.

    States print "4/24/25, 7/3/25, 7/18/25" or "12/31/2026-1/31/2027" for a
    layoff in waves; the first is when it begins.
    """
    if v is None or v == "":
        return None
    if isinstance(v, (dt.date, dt.datetime, int, float)):
        return parse_date(v)
    for tok in DATE_TOKEN.findall(str(v)):
        d = parse_date(tok)
        if d:
            return d
    return None


def parse_month(month, year) -> str | None:
    """'October' and 2026 to '2026-10'. For a state that prints only the month.

    A month is stored as the month: a day would be invented. The site prints
    it as "October 2026" and sorts it ahead of that month's dated notices.
    """
    if month is None or year is None:
        return None
    if isinstance(month, (dt.date, dt.datetime)):
        return f"{month.year:04d}-{month.month:02d}"
    key = str(month).strip().lower().rstrip(".")
    n = MONTHS.get(key)
    if not n:
        return None
    try:
        y = int(str(year).strip())
    except ValueError:
        return None
    return f"{y:04d}-{n:02d}"


def is_month_only(d: str | None) -> bool:
    return bool(d) and len(d) == 7


def as_of_date(d: str) -> str:
    """A notice date the freshness check can age: a month reads as its first day.

    The first day is the earliest the notice could have been posted, so a
    month-only state never reads fresher than it is.
    """
    return f"{d}-01" if is_month_only(d) else d


def parse_int(v) -> int | None:
    """A worker count, or None when the cell holds no single number.

    "1,114", " 52 " and "52 (estimated)" read as numbers; "50-75", "TBD" and an
    empty cell don't. None means the state printed no count, never zero.
    """
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v == int(v) else None
    s = str(v).replace("\xa0", " ").replace(",", "").strip()
    if not s:
        return None
    if re.fullmatch(r"\d+(\.0+)?", s):
        return int(float(s))
    groups = re.findall(r"\d+", s)
    return int(groups[0]) if len(groups) == 1 else None


def clean(v) -> str:
    """A name or place as printed, with whitespace collapsed."""
    if v is None:
        return ""
    return re.sub(r"\s+", " ", html.unescape(str(v)).replace("\xa0", " ")).strip()


def join_site(*parts) -> str | None:
    """A display address from its parts, skipping blanks: '1 Main St, Springfield, IL 62701'."""
    out = [clean(p) for p in parts if p is not None and clean(p)]
    return ", ".join(out) or None


# --- rows -------------------------------------------------------------------

def notice(state: str, notice_date: str, company, *, source_url: str, effective_date: str | None = None,
           kind=None, employees: int | None = None, county=None, industry=None, site=None,
           extra="") -> dict:
    """One row in the warn_notices shape. `extra` separates one company's
    notices filed the same day (an address, a city, the state's own record id)
    and is folded into the id by assign_ids; it is never stored."""
    return {
        "state": state,
        "notice_date": notice_date,
        "effective_date": effective_date,
        "company": clean(company),
        "kind": clean(kind) or None,
        "employees": employees,
        "county": clean(county) or None,
        "industry": clean(industry) or None,
        "source_url": source_url,
        "site": clean(site) or None,
        "_extra": clean(extra),
    }


def assign_ids(rows: list[dict]) -> list[dict]:
    """A stable id per notice.

    One employer can file several notices on one day for several sites, and a
    report can print identical rows outright; the parser's `_extra` field (an
    address, a city, an index) separates the first and a sequence number the
    second. Row order is stable between downloads of the same report, so the
    ids are stable too.
    """
    seen: dict[str, int] = {}
    for row in rows:
        # `employees` is deliberately not in the identity: a state revises a
        # worker count, and with the count in the hash a revision would become a
        # second row instead of an update. Exact duplicates are still separated
        # by the sequence number below, and `write()` compares the count, so a
        # revision is picked up.
        base = "|".join([row["state"], row["notice_date"], row["company"], row["effective_date"] or "", row["county"] or "", row.pop("_extra", "") or ""])
        n = seen.get(base, 0)
        seen[base] = n + 1
        row["id"] = hashlib.sha1(f"{base}|{n}".encode()).hexdigest()[:16]
    return rows


def finish(rows: list[dict], since: str | None) -> list[dict]:
    """Drop rows with no company or no notice date and those before `since`,
    then give the rest their ids. Every new state's parser ends here."""
    kept = [r for r in rows if r["company"] and r["notice_date"]
            and (since is None or as_of_date(r["notice_date"]) >= since)]
    return assign_ids(kept)


# --- fetching -----------------------------------------------------------------

# One retry after lib_gov_data's first backoff: a state site that answers 403 or
# 503 twice is down or refusing, and the run moves on to the next state.
FETCH_ATTEMPTS = 2
# A dropped connection is asked again this many times, a few seconds apart:
# tn.gov reset the first two connections on Oct 25 2026 and answered the third.
RESET_ATTEMPTS = 3
RESET_WAIT_S = 5


def get(url: str, referer: str | None = None) -> bytes:
    """GET through lib_gov_data.fetch: robots.txt first, the browser header set, a backoff."""
    for attempt in range(1, RESET_ATTEMPTS + 1):
        try:
            return gov_fetch(url, referer=referer, attempts=FETCH_ATTEMPTS)
        except urllib.error.HTTPError:
            raise
        except (urllib.error.URLError, ConnectionError, TimeoutError) as exc:
            if attempt == RESET_ATTEMPTS:
                raise
            print(f"  {type(exc).__name__} on {url[:80]} (attempt {attempt}/{RESET_ATTEMPTS}); waiting {RESET_WAIT_S}s",
                  flush=True)
            time.sleep(RESET_WAIT_S)
    raise RuntimeError("unreachable")


def get_text(url: str, referer: str | None = None) -> str:
    return decode(get(url, referer))


def decode(b: bytes) -> str:
    """UTF-8, falling back to Windows-1252 (Massachusetts's CSV is cp1252: 'Labouré')."""
    try:
        return b.decode("utf-8-sig")
    except UnicodeDecodeError:
        return b.decode("cp1252", "replace")


def post(url: str, data: bytes, headers: dict | None = None, referer: str | None = None,
         opener: urllib.request.OpenerDirector | None = None) -> bytes:
    """POST, after the same robots.txt check `get()` makes."""
    if not robots_allowed(url):
        raise RobotsDisallowed(f"{url} is a path the host's robots.txt asks automated clients not to read")
    h = dict(BROWSER_HEADERS)
    h.update({"Sec-Fetch-Mode": "cors", "Sec-Fetch-Dest": "empty", "Sec-Fetch-Site": "same-origin"})
    h.pop("Sec-Fetch-User", None)
    h.pop("Upgrade-Insecure-Requests", None)
    if referer:
        h["Referer"] = referer
    h.update(headers or {})
    req = urllib.request.Request(url, data=data, headers=h, method="POST")
    open_ = opener.open if opener else urllib.request.urlopen
    with open_(req, timeout=120) as r:
        return r.read()


def session() -> urllib.request.OpenerDirector:
    """An opener that keeps cookies, for a site whose form ties its token to a cookie."""
    import http.cookiejar
    return urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))


def get_with(opener: urllib.request.OpenerDirector, url: str, referer: str | None = None) -> bytes:
    """GET through a cookie-keeping opener, with the same robots.txt check."""
    if not robots_allowed(url):
        raise RobotsDisallowed(f"{url} is a path the host's robots.txt asks automated clients not to read")
    h = dict(BROWSER_HEADERS)
    if referer:
        h["Referer"] = referer
    with opener.open(urllib.request.Request(url, headers=h), timeout=120) as r:
        return r.read()


# Network trouble a run can't help: the state's site is down, slow or refusing.
# The state is named in the run's note and its own freshness row is the alarm
# (each has a budget), so one host's bad day doesn't turn the whole run red.
TRANSIENT_STATUS = {403, 408, 429, 500, 502, 503, 504, 520, 521, 522, 523, 524}


def is_unreachable(exc: BaseException) -> bool:
    import http.client
    import ssl
    if isinstance(exc, urllib.error.HTTPError):
        return exc.code in TRANSIENT_STATUS
    return isinstance(exc, (urllib.error.URLError, TimeoutError, ConnectionError, ssl.SSLError,
                            http.client.HTTPException))


# --- pages ----------------------------------------------------------------------

class Cell(NamedTuple):
    text: str
    links: tuple[str, ...]
    header: bool
    # The cell's text split where the page breaks it (<br>, <p>, <div>, <li>):
    # Florida prints the company, then its address, in one cell.
    lines: tuple[str, ...] = ()


BREAK = "\x00"


class _TableParser(HTMLParser):
    """Every <table> as rows of cells, nested tables kept apart."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[Cell]]] = []
        self._frames: list[dict] = []

    def _close_cell(self, f: dict) -> None:
        if f["cell"] is not None and f["row"] is not None:
            raw = "".join(f["cell"]["text"]).replace("\xa0", " ")
            text = re.sub(r"\s+", " ", raw.replace(BREAK, " ")).strip()
            lines = tuple(x for x in (re.sub(r"\s+", " ", part).strip() for part in raw.split(BREAK)) if x)
            f["row"].append(Cell(text, tuple(f["cell"]["links"]), f["cell"]["th"], lines))
        f["cell"] = None

    def _close_row(self, f: dict) -> None:
        self._close_cell(f)
        if f["row"]:
            f["rows"].append(f["row"])
        f["row"] = None

    def handle_starttag(self, tag, attrs):
        f = self._frames[-1] if self._frames else None
        if tag == "table":
            self._frames.append({"rows": [], "row": None, "cell": None})
        elif f is None:
            return
        elif tag == "tr":
            self._close_row(f)
            f["row"] = []
        elif tag in ("td", "th"):
            self._close_cell(f)
            if f["row"] is None:
                f["row"] = []
            f["cell"] = {"text": [], "links": [], "th": tag == "th"}
        elif f["cell"] is not None:
            if tag == "a":
                href = dict(attrs).get("href")
                if href:
                    f["cell"]["links"].append(href)
            elif tag in ("br", "p", "div", "li"):
                f["cell"]["text"].append(BREAK)

    def handle_endtag(self, tag):
        if not self._frames:
            return
        f = self._frames[-1]
        if tag in ("br", "p", "div", "li") and f["cell"] is not None:
            # Florida writes </br>; an end tag breaks the line as well.
            f["cell"]["text"].append(BREAK)
        elif tag in ("td", "th"):
            self._close_cell(f)
        elif tag == "tr":
            self._close_row(f)
        elif tag == "table":
            self._close_row(f)
            self.tables.append(self._frames.pop()["rows"])

    def handle_data(self, data):
        if self._frames and self._frames[-1]["cell"] is not None:
            self._frames[-1]["cell"]["text"].append(data)


def tables(page: str) -> list[list[list[Cell]]]:
    p = _TableParser()
    p.feed(page)
    p.close()
    while p._frames:  # an unclosed table at the end of the page
        f = p._frames.pop()
        p._close_row(f)
        p.tables.append(f["rows"])
    return p.tables


def table_with(page: str, *needles: str) -> tuple[list[str], list[list[Cell]]] | None:
    """The first table whose header row names every needle, as (header, body rows).

    The header is the first row; matched case-insensitively.
    """
    for t in tables(page):
        if not t:
            continue
        head = [c.text.lower() for c in t[0]]
        if all(any(n in h for h in head) for n in needles):
            return head, t[1:]
    return None


def col(header: list[str], *needles: str, exclude: tuple[str, ...] = ()) -> int | None:
    """The first header naming every needle (lower case), or None."""
    for i, h in enumerate(header):
        h = h.lower()
        if all(n in h for n in needles) and not any(x in h for x in exclude):
            return i
    return None


def cell(row, i: int | None):
    """A row's value at `i`, or None past its end or when `i` is None."""
    if i is None or i >= len(row):
        return None
    v = row[i]
    return v.text if isinstance(v, Cell) else v


def links(page: str, base: str) -> list[tuple[str, str]]:
    """Every <a href> on the page as (absolute url, link text)."""
    out = []
    for m in re.finditer(r'<a\b[^>]*?href\s*=\s*["\']([^"\']+)["\'][^>]*>(.*?)</a>', page, re.S | re.I):
        href = html.unescape(m[1]).strip()
        text = clean(re.sub(r"<[^>]+>", " ", m[2]))
        out.append((urllib.parse.urljoin(base, href), text))
    return out


def text_of(fragment: str) -> str:
    return clean(re.sub(r"<[^>]+>", " ", re.sub(r"<(script|style)\b.*?</\1>", " ", fragment, flags=re.S | re.I)))


# --- files ----------------------------------------------------------------------

def xlsx_sheets(data: bytes) -> dict[str, list[tuple]]:
    """Every sheet's rows as tuples, values as stored (dates as datetimes)."""
    import openpyxl  # installed by the workflow step; not in the standard library
    import warnings

    if not data.startswith(b"PK"):
        raise ValueError("not a spreadsheet: the state answered with something else")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return {name: list(wb[name].iter_rows(values_only=True)) for name in wb.sheetnames}


def find_header(rows: list, *needles: str) -> int | None:
    """The index of the first row naming every needle (lower case)."""
    for i, r in enumerate(rows):
        cells = [str(c).lower() for c in r if c is not None]
        if cells and all(any(n in c for c in cells) for n in needles):
            return i
    return None


def header_of(row) -> list[str]:
    return [re.sub(r"\s+", " ", str(c or "")).strip().lower() for c in row]


def csv_rows(data: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(decode(data))))


def pdf_tables(data: bytes) -> list[list[list[str | None]]]:
    """Every table pdfplumber finds, page by page.

    pdfplumber is the PDF reader the debarment and annual-limits loaders
    already use; the workflow installs it beside openpyxl. Imported here so a
    test of a non-PDF state needs no PDF library.
    """
    import pdfplumber

    if not data.startswith(b"%PDF"):
        raise ValueError("not a PDF: the state answered with something else")
    out = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            out.extend(page.extract_tables())
    return out
