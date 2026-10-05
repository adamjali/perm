#!/usr/bin/env python3
"""The firm email-domain table a page claim is checked against.

Runs build_firm_domains against a generated workbook and in-memory SQLite:

1. only the DOMAIN of an address is read, and malformed cells are dropped;
   styled empty rows past the last case are counted as blank, not as rows;
2. personal and internet-provider mail never counts as a firm's domain;
3. columns resolve by DOL's name, and a file with neither column refuses;
4. a firm's spelling lands on its page by program_key, a spelling with no page is skipped;
5. a domain DOL prints beside more than SHARED_LIMIT firms is dropped;
6. writes are diffs per program: a second identical run writes nothing, a
   vanished pair is removed, and another program's rows are untouched.

Run:  python3 scripts/test_firm_domains.py
"""
from __future__ import annotations

import pathlib
import sys
import tempfile
from collections import Counter

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build_firm_domains as b  # noqa: E402
from lib_sqlite_shim import SqliteTurso  # noqa: E402

FAILS: list[str] = []


def check(label: str, ok: bool) -> None:
    print(("PASS " if ok else "FAIL ") + label)
    if not ok:
        FAILS.append(label)


# 1. domains
check("domain is lowercased and the local part dropped", b.domain_of(" Jane.Doe@Fragomen.COM ") == "fragomen.com")
check("www. is not part of a mail domain", b.domain_of("a@www.firm-law.com") == "firm-law.com")
check("two addresses in one cell keep the first", b.domain_of("a@one.com; b@two.com") == "one.com")
check("angle brackets and a trailing dot are stripped", b.domain_of("<a@firm.com.>") == "firm.com")
check("no @ is no domain", b.domain_of("firm.com") is None)
check("a dotless host is no domain", b.domain_of("a@localhost") is None)
check("an empty cell is no domain", b.domain_of("") is None and b.domain_of(None) is None)

# 2. personal mail
for d in ("gmail.com", "yahoo.co.uk", "hotmail.fr", "comcast.net", "outlook.com", "proton.me"):
    check(f"{d} is personal", b.is_personal(d))
for d in ("fragomen.com", "bal.com", "smallfirmlaw.com"):
    check(f"{d} is a firm's", not b.is_personal(d))


# 3. reading a workbook
def workbook(rows: list[list[str]]) -> str:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    path = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False).name
    wb.save(path)
    return path


perm = workbook([
    ["CASE_NUMBER", "ATTY_AG_LAW_FIRM_NAME", "ATTY_AG_EMAIL", "EMP_POC_EMAIL"],
    ["G-100-1", "Fragomen, Del Rey, Bernsen & Loewy LLP", "a@fragomen.com", "hr@acme.com"],
    ["G-100-2", "FRAGOMEN DEL REY BERNSEN & LOEWY, LLP", "b@Fragomen.com", "hr@acme.com"],
    ["G-100-3", "Fragomen, Del Rey, Bernsen & Loewy LLP", "c@gmail.com", ""],
    ["G-100-4", "Small Firm PLLC", "x@smallfirmlaw.com", ""],
    ["G-100-5", "", "y@nofirm.com", ""],
    ["G-100-6", "Small Firm PLLC", "", ""],
])
pairs, firms, stats = b.read_pairs(perm, "perm")
frag = b.program_key("Fragomen, Del Rey, Bernsen & Loewy LLP")
check("two spellings of one firm share a key", b.program_key("FRAGOMEN DEL REY BERNSEN & LOEWY, LLP") == frag)
check("a firm's domain is counted per filing", pairs[(frag, "fragomen.com")] == 2)
check("gmail is never a pair", not any(d == "gmail.com" for _k, d in pairs))
check("a personal address still counts as an emailed filing", firms[frag] == 3)
check("a row with no firm is skipped", stats.get("no_firm") == 1 and not any(d == "nofirm.com" for _k, d in pairs))
check("a row with no email is skipped", stats.get("no_email") == 1)
check("the employer's contact column is never read", not any(d == "acme.com" for _k, d in pairs))

check("every case row is counted as a row", stats.get("rows") == 6)


def styled_blanks(path: str, n: int) -> str:
    """The same workbook with n styled, empty rows after the last case, as DOL's sheets carry."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(path)
    ws = wb.active
    last = ws.max_row
    for r in range(last + 1, last + 1 + n):
        ws.cell(row=r, column=1).font = Font(bold=True)
    wb.save(path)
    return path


_p, _f, blank_stats = b.read_pairs(styled_blanks(perm, 5), "perm")
check("a styled empty row is counted as blank, not as a row without a firm",
      blank_stats.get("blank") == 5 and blank_stats.get("rows") == 6 and blank_stats.get("no_firm") == 1)
check("blank rows change no pair", _p == pairs)

pw = workbook([
    ["CASE_NUMBER", "LAWFIRM_NAME_BUSINESS_NAME", "AGENT_ATTORNEY_EMAIL_ADDRESS"],
    ["P-100-1", "Small Firm PLLC", "z@smallfirmlaw.com"],
])
pw_pairs, _f, _s = b.read_pairs(pw, "pw")
check("the wage-request layout resolves by name", pw_pairs[(b.program_key("Small Firm PLLC"), "smallfirmlaw.com")] == 1)

bad = workbook([["CASE_NUMBER", "SOMETHING_ELSE"], ["G-100-1", "x"]])
try:
    b.read_pairs(bad, "perm")
    refused = False
except SystemExit:
    refused = True
check("a file with neither column refuses instead of reading nothing", refused)

# 4 and 5. placing on pages
pages = {
    "fragomen-del-rey-bernsen-loewy-llp": ("Fragomen, Del Rey, Bernsen & Loewy LLP",
                                           "FRAGOMEN DEL REY BERNSEN LOEWY", 48000),
    "small-firm-pllc": ("Small Firm PLLC", "SMALL FIRM", 12),
}
rows, pstats = b.plan("perm", pairs, firms, pages)
got = {(r[0], r[1]): (r[3], r[4]) for r in rows}
check("Fragomen's domain lands on its page with its counts",
      got.get(("fragomen-del-rey-bernsen-loewy-llp", "fragomen.com")) == (2, 3))
check("the small firm's domain lands on its page; a filing with no email isn't an emailed one",
      got.get(("small-firm-pllc", "smallfirmlaw.com")) == (1, 1))

orphan_pairs = Counter({(b.program_key("No Page Partners LLP"), "nopage.com"): 4})
orphan_firms = Counter({b.program_key("No Page Partners LLP"): 4})
orows, _ = b.plan("perm", orphan_pairs, orphan_firms, pages)
check("a firm with no page is skipped", orows == [])

many_pages = {f"firm-{i}": (f"Firm Number {i} LLP", f"FIRM NUMBER {i}", 10) for i in range(5)}
shared_pairs = Counter({(b.program_key(f"Firm Number {i} LLP"), "sharedservice.com"): 3 for i in range(5)})
shared_pairs[(b.program_key("Firm Number 0 LLP"), "firmzero.com")] = 3
shared_firms = Counter({b.program_key(f"Firm Number {i} LLP"): 3 for i in range(5)})
srows, sstats = b.plan("lca", shared_pairs, shared_firms, many_pages)
check("a domain beside more than SHARED_LIMIT firms is dropped",
      not any(r[1] == "sharedservice.com" for r in srows) and sstats["shared_domains_dropped"] == 1)
check("a firm's own domain survives beside a shared one", any(r[:2] == ["firm-0", "firmzero.com"] for r in srows))

# 6. writes
db = SqliteTurso()
db.url = "sqlite://memory"
for ddl in b.DDL:
    db.execute(ddl)
b.WRITE_PAUSE_S = 0
w1, g1 = b.write_program(db, "perm", rows)
check("the first write stores every row", (w1, g1) == (len(rows), 0))
b.write_program(db, "pw", [["small-firm-pllc", "smallfirmlaw.com", "pw", 1, 1]])
w2, g2 = b.write_program(db, "perm", rows)
check("an identical second run writes nothing", (w2, g2) == (0, 0))
w3, g3 = b.write_program(db, "perm", [r for r in rows if r[0] != "small-firm-pllc"])
left = b.rows_of(db.execute("SELECT page_slug, program FROM firm_email_domains ORDER BY page_slug, program"))
check("a vanished pair is removed", (w3, g3) == (0, 1))
check("another program's rows are untouched",
      [list(r) for r in left] == [["fragomen-del-rey-bernsen-loewy-llp", "perm"], ["small-firm-pllc", "pw"]])
changed = [[r[0], r[1], r[2], r[3] + 1, r[4]] for r in rows if r[0] != "small-firm-pllc"]
w4, _ = b.write_program(db, "perm", changed)
check("a changed count is rewritten", w4 == 1)

if FAILS:
    print(f"\n{len(FAILS)} FAILED")
    sys.exit(1)
print("\nall passed")
